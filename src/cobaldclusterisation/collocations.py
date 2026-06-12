"""Collocation tables for CoBaLD corpus inspection."""

from __future__ import annotations

from collections import Counter
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass
from itertools import combinations
from pathlib import Path
import pickle
from typing import Literal

import pandas as pd

from .data import Sentence, Token, load_corpus


TokenBasis = Literal["lemma", "form"]
OccurrenceIdIndex = dict[tuple[str, int, str], list[int]]

DEFAULT_TOKEN_BASES: tuple[TokenBasis, ...] = ("lemma", "form")
DEFAULT_NGRAM_SIZES: tuple[int, ...] = (2, 3)
DEFAULT_SORT_BY = "pmi"
DEFAULT_MIN_FREQ = 3
DEFAULT_CLUSTER_MIN_FREQ = 5
DEFAULT_CLUSTER_ARTIFACT_PATH = (
    "/content/drive/MyDrive/cobald_outputs/results/"
    "100,200,300,400,512,565cl_rubert_tiny2_Minibatch_Kmeans_artifacts.pkl"
)
DEFAULT_CLUSTER_OUTPUT_SUFFIX = "collocation_cluster_overlap.xlsx"
INVALID_SEMCLASS_VALUES = {"", "_", "nan", "None"}
VALID_SORT_COLUMNS = {
    "frequency",
    "raw_freq",
    "pmi",
    "likelihood_ratio",
    "student_t",
    "chi_sq",
}


@dataclass(frozen=True, slots=True)
class CollocationConfig:
    """Configuration for CoBaLD collocation extraction."""

    data_dir: str | Path = "CobaldRus"
    splits: Sequence[str] = ("train", "dev")
    token_bases: Sequence[TokenBasis] = DEFAULT_TOKEN_BASES
    ngram_sizes: Sequence[int] = DEFAULT_NGRAM_SIZES
    min_freq: int = DEFAULT_MIN_FREQ
    sort_by: str = DEFAULT_SORT_BY
    max_rows: int | None = None
    lowercase: bool = True
    output_path: str | Path = "outputs/collocations/cobald_collocations.xlsx"


@dataclass(frozen=True, slots=True)
class CollocationClusterConfig:
    """Configuration for collocation comparison against clustering artifacts."""

    artifact_path: str | Path = DEFAULT_CLUSTER_ARTIFACT_PATH
    data_dir: str | Path = "CobaldRus"
    splits: Sequence[str] = ("train", "dev")
    token_bases: Sequence[TokenBasis] = DEFAULT_TOKEN_BASES
    ngram_sizes: Sequence[int] = DEFAULT_NGRAM_SIZES
    min_freq: int = DEFAULT_CLUSTER_MIN_FREQ
    sort_by: str = DEFAULT_SORT_BY
    max_rows: int | None = None
    lowercase: bool = True
    output_path: str | Path | None = None
    max_examples_per_run: int = 300
    create_plots: bool = True


def _require_nltk() -> tuple[object, object, object, object]:
    try:
        from nltk.collocations import BigramCollocationFinder, TrigramCollocationFinder
        from nltk.metrics import BigramAssocMeasures, TrigramAssocMeasures
    except ImportError as exc:
        raise ImportError(
            "Collocation extraction requires nltk. Install it in Colab with: "
            'pip install -e ".[collocations]" openpyxl'
        ) from exc
    return (
        BigramCollocationFinder,
        TrigramCollocationFinder,
        BigramAssocMeasures,
        TrigramAssocMeasures,
    )


def _validate_config(config: CollocationConfig) -> None:
    if config.min_freq < 1:
        raise ValueError("min_freq must be positive")
    if config.max_rows is not None and config.max_rows < 1:
        raise ValueError("max_rows must be positive when provided")
    if config.sort_by not in VALID_SORT_COLUMNS:
        valid = ", ".join(sorted(VALID_SORT_COLUMNS))
        raise ValueError(f"sort_by must be one of: {valid}")
    unsupported = sorted(set(config.ngram_sizes) - {2, 3})
    if unsupported:
        raise ValueError(f"Only 2-grams and 3-grams are supported, got: {unsupported}")
    unsupported_bases = sorted(set(config.token_bases) - {"lemma", "form"})
    if unsupported_bases:
        raise ValueError(
            f"token_bases must contain only 'lemma' and/or 'form', got: {unsupported_bases}"
        )


def _validate_cluster_config(config: CollocationClusterConfig) -> None:
    _validate_config(
        CollocationConfig(
            data_dir=config.data_dir,
            splits=config.splits,
            token_bases=config.token_bases,
            ngram_sizes=config.ngram_sizes,
            min_freq=config.min_freq,
            sort_by=config.sort_by,
            max_rows=config.max_rows,
            lowercase=config.lowercase,
        )
    )
    if config.max_examples_per_run < 0:
        raise ValueError("max_examples_per_run must be non-negative")


def token_text(token: Token, *, basis: TokenBasis, lowercase: bool = True) -> str | None:
    """Return normalized token text for one token and basis."""

    value = token.lemma if basis == "lemma" else token.form
    if value in {"", "_"}:
        return None
    return value.lower() if lowercase else value


def sentence_token_documents(
    sentences: Iterable[Sentence],
    *,
    basis: TokenBasis,
    lowercase: bool = True,
) -> list[list[str]]:
    """Return sentence-bounded token documents for NLTK collocation finders."""

    documents: list[list[str]] = []
    for sentence in sentences:
        document = [
            text
            for token in sentence.surface_tokens(include_punctuation=False)
            if (text := token_text(token, basis=basis, lowercase=lowercase)) is not None
        ]
        if document:
            documents.append(document)
    return documents


def _score_ngram(
    finder: object,
    ngram: tuple[str, ...],
    measures: object,
    metric_names: Sequence[str],
) -> dict[str, float | int | str]:
    score_ngram: Callable[..., float] = getattr(finder, "score_ngram")
    row: dict[str, float | int | str] = {
        "ngram": " ".join(ngram),
        "frequency": int(finder.ngram_fd[ngram]),
    }
    for metric_name in metric_names:
        metric = getattr(measures, metric_name)
        row[metric_name] = float(score_ngram(metric, *ngram))
    return row


def collocation_table_from_documents(
    documents: Sequence[Sequence[str]],
    *,
    ngram_size: int,
    min_freq: int = DEFAULT_MIN_FREQ,
    sort_by: str = DEFAULT_SORT_BY,
    max_rows: int | None = None,
) -> pd.DataFrame:
    """Build a scored collocation table from sentence-bounded token documents."""

    if ngram_size not in {2, 3}:
        raise ValueError("ngram_size must be 2 or 3")
    if min_freq < 1:
        raise ValueError("min_freq must be positive")
    if sort_by not in VALID_SORT_COLUMNS:
        valid = ", ".join(sorted(VALID_SORT_COLUMNS))
        raise ValueError(f"sort_by must be one of: {valid}")

    (
        BigramCollocationFinder,
        TrigramCollocationFinder,
        BigramAssocMeasures,
        TrigramAssocMeasures,
    ) = _require_nltk()

    if ngram_size == 2:
        finder = BigramCollocationFinder.from_documents(documents)
        measures = BigramAssocMeasures()
    else:
        finder = TrigramCollocationFinder.from_documents(documents)
        measures = TrigramAssocMeasures()
    finder.apply_freq_filter(min_freq)

    metric_names = ("raw_freq", "pmi", "likelihood_ratio", "student_t", "chi_sq")
    rows = [
        _score_ngram(finder, ngram, measures, metric_names)
        for ngram in finder.ngram_fd
    ]
    table = pd.DataFrame(rows)
    if table.empty:
        return pd.DataFrame(
            columns=[
                "ngram",
                "frequency",
                "raw_freq",
                "pmi",
                "likelihood_ratio",
                "student_t",
                "chi_sq",
            ]
        )

    table = table.sort_values(
        by=[sort_by, "frequency", "ngram"],
        ascending=[False, False, True],
        kind="mergesort",
    ).reset_index(drop=True)
    if max_rows is not None:
        table = table.head(max_rows).reset_index(drop=True)
    return table


def build_collocation_tables(
    sentences: Sequence[Sentence],
    *,
    token_bases: Sequence[TokenBasis] = DEFAULT_TOKEN_BASES,
    ngram_sizes: Sequence[int] = DEFAULT_NGRAM_SIZES,
    min_freq: int = DEFAULT_MIN_FREQ,
    sort_by: str = DEFAULT_SORT_BY,
    max_rows: int | None = None,
    lowercase: bool = True,
) -> dict[str, pd.DataFrame]:
    """Build all requested collocation tables for loaded CoBaLD sentences."""

    config = CollocationConfig(
        token_bases=token_bases,
        ngram_sizes=ngram_sizes,
        min_freq=min_freq,
        sort_by=sort_by,
        max_rows=max_rows,
        lowercase=lowercase,
    )
    _validate_config(config)

    tables: dict[str, pd.DataFrame] = {}
    for basis in token_bases:
        documents = sentence_token_documents(
            sentences,
            basis=basis,
            lowercase=lowercase,
        )
        for ngram_size in ngram_sizes:
            name = f"{basis}_{'bigrams' if ngram_size == 2 else 'trigrams'}"
            table = collocation_table_from_documents(
                documents,
                ngram_size=ngram_size,
                min_freq=min_freq,
                sort_by=sort_by,
                max_rows=max_rows,
            )
            table.insert(0, "token_basis", basis)
            table.insert(1, "ngram_size", ngram_size)
            tables[name] = table
    return tables


def _metadata_table(
    *,
    config: CollocationConfig,
    sentences: Sequence[Sentence],
    tables: dict[str, pd.DataFrame],
) -> pd.DataFrame:
    token_counts = {
        basis: sum(
            len(document)
            for document in sentence_token_documents(
                sentences,
                basis=basis,
                lowercase=config.lowercase,
            )
        )
        for basis in config.token_bases
    }
    rows: list[dict[str, object]] = [
        {"key": "data_dir", "value": str(config.data_dir)},
        {"key": "splits", "value": ",".join(config.splits)},
        {"key": "sentence_count", "value": len(sentences)},
        {"key": "token_bases", "value": ",".join(config.token_bases)},
        {"key": "ngram_sizes", "value": ",".join(str(n) for n in config.ngram_sizes)},
        {"key": "min_freq", "value": config.min_freq},
        {"key": "sort_by", "value": config.sort_by},
        {"key": "max_rows", "value": "" if config.max_rows is None else config.max_rows},
        {"key": "lowercase", "value": config.lowercase},
        {"key": "token_filter", "value": "surface non-punctuation tokens"},
    ]
    rows.extend(
        {"key": f"{basis}_token_count", "value": count}
        for basis, count in token_counts.items()
    )
    rows.extend(
        {"key": f"{name}_rows", "value": len(table)}
        for name, table in tables.items()
    )
    return pd.DataFrame(rows)


def write_collocation_excel(
    tables: dict[str, pd.DataFrame],
    output_path: str | Path,
    *,
    metadata: pd.DataFrame | None = None,
) -> str:
    """Write collocation tables to an Excel workbook."""

    try:
        from openpyxl.styles import Alignment
    except ImportError as exc:
        raise ImportError(
            "Excel export requires openpyxl. Install it in Colab with: "
            "pip install openpyxl"
        ) from exc

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, table in tables.items():
            table.to_excel(writer, sheet_name=sheet_name[:31], index=False)
            worksheet = writer.sheets[sheet_name[:31]]
            for row in worksheet.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            worksheet.freeze_panes = "A2"
        if metadata is not None:
            metadata.to_excel(writer, sheet_name="metadata", index=False)
            worksheet = writer.sheets["metadata"]
            for row in worksheet.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            worksheet.freeze_panes = "A2"
    return str(output_path)


def load_clustering_artifact(path: str | Path) -> dict[str, object]:
    """Load a clustering ``*_artifacts.pkl`` file for collocation analysis."""

    with Path(path).open("rb") as handle:
        artifact = pickle.load(handle)
    if not isinstance(artifact, dict) or not isinstance(artifact.get("runs"), list):
        raise ValueError("artifact must be a dictionary containing a 'runs' list")
    return artifact


def _run_label(run_index: int, run: dict[str, object]) -> str:
    config = run.get("config", {})
    if not isinstance(config, dict):
        return f"run{run_index}"
    algorithm = str(config.get("algorithm", "run"))
    if "n_clusters" in config and config.get("n_clusters") is not None:
        return f"k{int(config['n_clusters'])}"
    if algorithm in {"knn_leiden", "knn_louvain"}:
        return f"nn{config.get('graph_n_neighbors', run_index)}"
    if algorithm == "hdbscan":
        return f"min{config.get('min_cluster_size', run_index)}"
    return f"run{run_index}"


def _sheet_name(base: str, used: set[str]) -> str:
    candidate = base[:31]
    if candidate not in used:
        used.add(candidate)
        return candidate
    index = 2
    while True:
        suffix = f"_{index}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        if candidate not in used:
            used.add(candidate)
            return candidate
        index += 1


def _mean_or_nan(numerator: float, denominator: int) -> float:
    return float(numerator / denominator) if denominator else float("nan")


def _format_counter(counter: Counter[object], *, max_items: int = 12) -> str:
    return "\n".join(f"{key}: {value}" for key, value in counter.most_common(max_items))


def _valid_semclass(value: object) -> bool:
    return str(value) not in INVALID_SEMCLASS_VALUES


def _row_token_text(
    row: pd.Series,
    *,
    basis: TokenBasis,
    lowercase: bool,
) -> str | None:
    value = row.get("LEMMA" if basis == "lemma" else "FORM")
    return _normalize_token_text_value(value, lowercase=lowercase)


def _normalize_token_text_value(
    value: object,
    *,
    lowercase: bool,
) -> str | None:
    if value is None:
        return None
    text = str(value)
    if text in {"", "_"}:
        return None
    return text.lower() if lowercase else text


def _collocation_occurrences(
    sentences: Sequence[Sentence],
    tables: dict[str, pd.DataFrame],
    *,
    token_bases: Sequence[TokenBasis],
    ngram_sizes: Sequence[int],
    lowercase: bool,
) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    occurrence_id = 0
    selected: dict[tuple[str, int], set[str]] = {}
    for basis in token_bases:
        for ngram_size in ngram_sizes:
            name = f"{basis}_{'bigrams' if ngram_size == 2 else 'trigrams'}"
            table = tables.get(name, pd.DataFrame())
            selected[(basis, ngram_size)] = (
                set(table["ngram"].astype(str)) if "ngram" in table else set()
            )

    for basis in token_bases:
        for sentence in sentences:
            token_rows = [
                (token, text)
                for token in sentence.surface_tokens(include_punctuation=False)
                if (text := token_text(token, basis=basis, lowercase=lowercase))
                is not None
            ]
            if not token_rows:
                continue
            tokens = [item[0] for item in token_rows]
            texts = [item[1] for item in token_rows]
            for ngram_size in ngram_sizes:
                accepted = selected.get((basis, ngram_size), set())
                if not accepted or len(texts) < ngram_size:
                    continue
                for start in range(0, len(texts) - ngram_size + 1):
                    parts = tuple(texts[start : start + ngram_size])
                    ngram = " ".join(parts)
                    if ngram not in accepted:
                        continue
                    occurrence_tokens = tokens[start : start + ngram_size]
                    row: dict[str, object] = {
                        "occurrence_id": occurrence_id,
                        "token_basis": basis,
                        "ngram_size": ngram_size,
                        "ngram": ngram,
                        "split": sentence.split,
                        "sentence_index": sentence.index,
                        "sentence_id": sentence.sent_id,
                        "context_text": sentence.text,
                        "_parts": parts,
                        "_token_keys": tuple(
                            (
                                token.split,
                                token.sentence_index,
                                token.token_index,
                            )
                            for token in occurrence_tokens
                        ),
                    }
                    for part_index, (part, token) in enumerate(
                        zip(parts, occurrence_tokens, strict=True),
                        start=1,
                    ):
                        row[f"part_{part_index}"] = part
                        row[f"part_{part_index}_FORM"] = token.form
                        row[f"part_{part_index}_LEMMA"] = token.lemma
                        row[f"part_{part_index}_SEMCLASS"] = token.semclass
                        row[f"part_{part_index}_token_index"] = token.token_index
                    rows.append(row)
                    occurrence_id += 1
    return pd.DataFrame(rows)


def _occurrence_id_index(occurrences: pd.DataFrame) -> OccurrenceIdIndex:
    if occurrences.empty:
        return {}
    required = {"token_basis", "ngram_size", "ngram", "occurrence_id"}
    missing = sorted(required - set(occurrences.columns))
    if missing:
        raise ValueError(f"occurrences is missing required columns: {missing}")

    grouped = occurrences.groupby(
        ["token_basis", "ngram_size", "ngram"],
        sort=False,
    )["occurrence_id"].agg(list)
    return {
        (str(basis), int(ngram_size), str(ngram)): [
            int(occurrence_id) for occurrence_id in occurrence_ids
        ]
        for (basis, ngram_size, ngram), occurrence_ids in grouped.items()
    }


def _annotated_frame(run: dict[str, object], *, run_index: int) -> pd.DataFrame:
    annotated = run.get("annotated_tokens")
    if not isinstance(annotated, pd.DataFrame):
        raise ValueError(f"run {run_index} has no annotated_tokens DataFrame")
    required = {"cluster", "FORM", "LEMMA", "SEMCLASS"}
    missing = sorted(required - set(annotated.columns))
    if missing:
        raise ValueError(f"run {run_index} annotated_tokens is missing: {missing}")
    return annotated.reset_index(drop=True).copy()


def _token_lookup(
    frame: pd.DataFrame,
) -> dict[tuple[object, object, object], dict[str, object]]:
    required = {"split", "sentence_index", "token_index"}
    missing = sorted(required - set(frame.columns))
    if missing:
        raise ValueError(
            "annotated_tokens must contain split, sentence_index, and token_index "
            f"for exact collocation matching; missing: {missing}"
        )
    lookup: dict[tuple[object, object, object], dict[str, object]] = {}
    columns = [
        column
        for column in (
            "split",
            "sentence_index",
            "token_index",
            "FORM",
            "LEMMA",
            "SEMCLASS",
            "cluster",
            "AUTO_SEMCLASS",
        )
        if column in frame.columns
    ]
    for row in frame[columns].to_dict("records"):
        key = (row.get("split"), row.get("sentence_index"), row.get("token_index"))
        lookup[key] = row
    return lookup


def _cluster_names(frame: pd.DataFrame) -> dict[int, str]:
    if "AUTO_SEMCLASS" not in frame:
        return {}
    names: dict[int, str] = {}
    for cluster, part in frame.groupby("cluster", sort=True):
        values = part["AUTO_SEMCLASS"].dropna().astype(str)
        if not values.empty:
            names[int(cluster)] = values.iloc[0]
    return names


def _format_cluster(cluster: int, names: dict[int, str]) -> str:
    name = names.get(cluster)
    return f"{cluster}:{name}" if name and name != str(cluster) else str(cluster)


def _shared_pair_summary(
    left: Counter[int],
    right: Counter[int],
    *,
    names: dict[int, str],
    max_items: int = 12,
) -> str:
    rows: list[tuple[int, int, int, int]] = []
    for cluster in set(left) & set(right):
        rows.append(
            (
                int(left[cluster]) * int(right[cluster]),
                cluster,
                left[cluster],
                right[cluster],
            )
        )
    rows.sort(reverse=True)
    return "\n".join(
        f"{_format_cluster(cluster, names)}: {left_count}+{right_count} pairs={pairs}"
        for pairs, cluster, left_count, right_count in rows[:max_items]
    )


def _shared_all_summary(
    counters: Sequence[Counter[int]],
    *,
    names: dict[int, str],
    max_items: int = 12,
) -> str:
    if not counters:
        return ""
    shared = set(counters[0])
    for counter in counters[1:]:
        shared &= set(counter)
    rows: list[tuple[int, int, tuple[int, ...]]] = []
    for cluster in shared:
        counts = tuple(int(counter[cluster]) for counter in counters)
        product = 1
        for count in counts:
            product *= count
        rows.append((product, int(cluster), counts))
    rows.sort(reverse=True)
    return "\n".join(
        f"{_format_cluster(cluster, names)}: counts={','.join(str(c) for c in counts)} tuples={product}"
        for product, cluster, counts in rows[:max_items]
    )


def _aggregate_maps(
    frame: pd.DataFrame,
    *,
    basis: TokenBasis,
    lowercase: bool,
) -> dict[str, dict[str, object]]:
    maps: dict[str, dict[str, object]] = {}
    token_column = "LEMMA" if basis == "lemma" else "FORM"
    for value, cluster, semclass in zip(
        frame[token_column].tolist(),
        frame["cluster"].tolist(),
        frame["SEMCLASS"].tolist(),
        strict=False,
    ):
        text = _normalize_token_text_value(value, lowercase=lowercase)
        if text is None:
            continue
        entry = maps.setdefault(
            text,
            {
                "token_count": 0,
                "clusters": Counter(),
                "semclasses": Counter(),
            },
        )
        entry["token_count"] = int(entry["token_count"]) + 1
        entry["clusters"][int(cluster)] += 1
        if _valid_semclass(semclass):
            entry["semclasses"][str(semclass)] += 1
    return maps


def _part_stats(
    parts: Sequence[str],
    aggregate: dict[str, dict[str, object]],
) -> dict[str, object]:
    row: dict[str, object] = {}
    for index, part in enumerate(parts, start=1):
        stats = aggregate.get(part, {})
        clusters = stats.get("clusters", Counter())
        semclasses = stats.get("semclasses", Counter())
        assert isinstance(clusters, Counter)
        assert isinstance(semclasses, Counter)
        dominant = semclasses.most_common(1)[0][0] if semclasses else ""
        row[f"part_{index}"] = part
        row[f"part_{index}_token_count"] = int(stats.get("token_count", 0))
        row[f"part_{index}_cluster_count"] = len(clusters)
        row[f"part_{index}_clusters"] = _format_counter(clusters)
        row[f"part_{index}_dominant_semclass"] = dominant
        row[f"part_{index}_semclasses"] = _format_counter(semclasses)
    return row


def _cluster_collocation_table(tables: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """Return lemma collocations for cluster inspection, dropping embedded bigrams."""

    frames: list[pd.DataFrame] = []
    for table_name in ("lemma_bigrams", "lemma_trigrams"):
        table = tables.get(table_name, pd.DataFrame())
        if table.empty:
            continue
        frame = table.copy()
        frame["source_sheet"] = table_name
        frame["source_rank"] = range(1, len(frame) + 1)
        frames.append(frame)
    if not frames:
        return pd.DataFrame()

    collocations = pd.concat(frames, ignore_index=True, sort=False)
    trigram_bigrams: set[str] = set()
    for ngram in collocations.loc[
        collocations["ngram_size"].astype(int) == 3,
        "ngram",
    ].astype(str):
        parts = tuple(ngram.split(" "))
        trigram_bigrams.update(
            " ".join(parts[index : index + 2])
            for index in range(0, len(parts) - 1)
        )
    if trigram_bigrams:
        collocations = collocations.loc[
            ~(
                (collocations["ngram_size"].astype(int) == 2)
                & collocations["ngram"].astype(str).isin(trigram_bigrams)
            )
        ].copy()
    if collocations.empty:
        return collocations
    return collocations.sort_values(
        by=["frequency", "ngram_size", "pmi", "ngram"],
        ascending=[False, False, False, True],
        na_position="last",
        kind="mergesort",
    ).reset_index(drop=True)


def _collocation_semclass_counters(
    occurrence_ids: Sequence[int],
    occurrence_infos: dict[int, list[dict[str, object]]],
    *,
    part_count: int,
) -> list[Counter[str]]:
    counters = [Counter() for _ in range(part_count)]
    for occurrence_id in occurrence_ids:
        infos = occurrence_infos.get(int(occurrence_id), [])
        for index in range(min(part_count, len(infos))):
            semclass = infos[index].get("SEMCLASS")
            if _valid_semclass(semclass):
                counters[index][str(semclass)] += 1
    return counters


def _format_semclass_summary(
    collocation_semclasses: Counter[str],
    all_semclasses: Counter[str],
) -> str:
    dominant = (
        collocation_semclasses.most_common(1)[0][0]
        if collocation_semclasses
        else ""
    )
    all_values = ", ".join(label for label, _ in all_semclasses.most_common())
    return "\n".join(
        line
        for line in (
            f"dominant in collocation: {dominant}" if dominant else "",
            f"all: {all_values}" if all_values else "",
        )
        if line
    )


def _shared_cluster_count(counters: Sequence[Counter[int]]) -> int:
    if not counters:
        return 0
    shared = set(counters[0])
    for counter in counters[1:]:
        shared &= set(counter)
    return len(shared)


def _simple_collocation_row(
    table_row: pd.Series,
    *,
    aggregate: dict[str, dict[str, object]],
    occurrence_ids: Sequence[int],
    occurrence_infos: dict[int, list[dict[str, object]]],
) -> dict[str, object]:
    ngram = str(table_row["ngram"])
    parts = tuple(ngram.split(" "))
    collocation_semclasses = _collocation_semclass_counters(
        occurrence_ids,
        occurrence_infos,
        part_count=len(parts),
    )
    cluster_counters: list[Counter[int]] = []
    row: dict[str, object] = {
        "collocation": ngram,
        "frequency": int(table_row["frequency"]),
    }
    cluster_count_lines: list[str] = []
    for index, part in enumerate(parts, start=1):
        stats = aggregate.get(part, {})
        all_semclasses = stats.get("semclasses", Counter())
        clusters = stats.get("clusters", Counter())
        assert isinstance(all_semclasses, Counter)
        assert isinstance(clusters, Counter)
        cluster_counters.append(clusters)
        row[f"part_{index}"] = part
        row[f"part_{index}_semclasses"] = _format_semclass_summary(
            collocation_semclasses[index - 1],
            all_semclasses,
        )
        cluster_count_lines.append(f"{part}: {len(clusters)}")
    for index in range(len(parts) + 1, 4):
        row[f"part_{index}"] = ""
        row[f"part_{index}_semclasses"] = ""
    row["cluster_counts"] = "\n".join(cluster_count_lines)
    row["intersecting_cluster_count"] = _shared_cluster_count(cluster_counters)
    return row


def _counter_product_total(counters: Sequence[Counter[int]]) -> int:
    total = 1
    for counter in counters:
        total *= sum(counter.values())
    return int(total)


def _counter_all_same_count(counters: Sequence[Counter[int]]) -> int:
    if not counters:
        return 0
    shared = set(counters[0])
    for counter in counters[1:]:
        shared &= set(counter)
    total = 0
    for cluster in shared:
        product = 1
        for counter in counters:
            product *= int(counter[cluster])
        total += product
    return total


def _pair_metric(
    left: Counter[int],
    right: Counter[int],
) -> tuple[int, int, float]:
    total = int(sum(left.values()) * sum(right.values()))
    same = int(sum(left[cluster] * right[cluster] for cluster in set(left) & set(right)))
    return total, same, _mean_or_nan(float(same), total)


def _exact_rows_for_run(
    occurrences: pd.DataFrame,
    frame: pd.DataFrame,
    *,
    max_examples: int,
) -> tuple[dict[int, list[dict[str, object]]], pd.DataFrame]:
    lookup = _token_lookup(frame)
    occurrence_infos: dict[int, list[dict[str, object]]] = {}
    example_rows: list[dict[str, object]] = []
    for _, occurrence in occurrences.iterrows():
        token_infos: list[dict[str, object]] = []
        for key in occurrence["_token_keys"]:
            token_row = lookup.get(key)
            if token_row is None:
                token_infos.append({"matched": False})
                continue
            token_infos.append(
                {
                    "matched": True,
                    "FORM": token_row.get("FORM"),
                    "LEMMA": token_row.get("LEMMA"),
                    "SEMCLASS": token_row.get("SEMCLASS"),
                    "cluster": int(token_row.get("cluster")),
                    "AUTO_SEMCLASS": token_row.get("AUTO_SEMCLASS"),
                }
            )
        occurrence_infos[int(occurrence["occurrence_id"])] = token_infos
        if len(example_rows) >= max_examples:
            continue
        example: dict[str, object] = {
            "token_basis": occurrence["token_basis"],
            "ngram_size": occurrence["ngram_size"],
            "ngram": occurrence["ngram"],
            "occurrence_id": occurrence["occurrence_id"],
            "split": occurrence.get("split"),
            "sentence_index": occurrence.get("sentence_index"),
            "sentence_id": occurrence.get("sentence_id"),
            "context_text": occurrence.get("context_text"),
        }
        for index, info in enumerate(token_infos, start=1):
            example[f"part_{index}"] = occurrence.get(f"part_{index}")
            example[f"part_{index}_FORM"] = info.get("FORM")
            example[f"part_{index}_LEMMA"] = info.get("LEMMA")
            example[f"part_{index}_SEMCLASS"] = info.get("SEMCLASS")
            example[f"part_{index}_cluster"] = info.get("cluster")
            example[f"part_{index}_AUTO_SEMCLASS"] = info.get("AUTO_SEMCLASS")
            example[f"part_{index}_matched"] = info.get("matched")
        matched_clusters = [
            info["cluster"] for info in token_infos if info.get("matched")
        ]
        example["all_parts_matched"] = len(matched_clusters) == len(token_infos)
        example["all_parts_same_cluster"] = (
            len(matched_clusters) == len(token_infos)
            and len(set(matched_clusters)) == 1
        )
        example_rows.append(example)
    return occurrence_infos, pd.DataFrame(example_rows)


def _exact_collocation_metrics(
    occurrence_ids: Sequence[int],
    occurrence_infos: dict[int, list[dict[str, object]]],
) -> dict[str, int | float]:
    pair_total = 0
    pair_same = 0
    all_total = 0
    all_same = 0
    for occurrence_id in occurrence_ids:
        infos = occurrence_infos.get(int(occurrence_id), [])
        clusters = [
            int(info["cluster"])
            for info in infos
            if info.get("matched") and info.get("cluster") is not None
        ]
        if len(clusters) == len(infos) and clusters:
            all_total += 1
            if len(set(clusters)) == 1:
                all_same += 1
        for left, right in combinations(infos, 2):
            if (
                left.get("matched")
                and right.get("matched")
                and left.get("cluster") is not None
                and right.get("cluster") is not None
            ):
                pair_total += 1
                if int(left["cluster"]) == int(right["cluster"]):
                    pair_same += 1
    return {
        "exact_occurrence_count": len(occurrence_ids),
        "exact_all_parts_observation_count": all_total,
        "exact_all_parts_same_cluster_count": all_same,
        "exact_all_parts_same_cluster_rate": _mean_or_nan(float(all_same), all_total),
        "exact_pair_observation_count": pair_total,
        "exact_same_cluster_pair_count": pair_same,
        "exact_same_cluster_pair_rate": _mean_or_nan(float(pair_same), pair_total),
    }


def _exact_pair_metrics(
    occurrence_ids: Sequence[int],
    occurrence_infos: dict[int, list[dict[str, object]]],
    *,
    left_index: int,
    right_index: int,
) -> tuple[int, int, float]:
    total = 0
    same = 0
    for occurrence_id in occurrence_ids:
        infos = occurrence_infos.get(int(occurrence_id), [])
        if max(left_index, right_index) >= len(infos):
            continue
        left = infos[left_index]
        right = infos[right_index]
        if (
            left.get("matched")
            and right.get("matched")
            and left.get("cluster") is not None
            and right.get("cluster") is not None
        ):
            total += 1
            if int(left["cluster"]) == int(right["cluster"]):
                same += 1
    return total, same, _mean_or_nan(float(same), total)


def analyze_collocation_cluster_run(
    tables: dict[str, pd.DataFrame],
    occurrences: pd.DataFrame,
    run: dict[str, object],
    *,
    run_index: int,
    lowercase: bool,
    max_examples: int,
    occurrence_id_index: Mapping[tuple[str, int, str], Sequence[int]] | None = None,
) -> dict[str, pd.DataFrame | dict[str, object]]:
    """Analyze one clustering run against already computed collocation tables."""

    frame = _annotated_frame(run, run_index=run_index)
    if occurrence_id_index is None:
        occurrence_id_index = _occurrence_id_index(occurrences)
    occurrence_infos, _ = _exact_rows_for_run(
        occurrences,
        frame,
        max_examples=0,
    )
    aggregate = _aggregate_maps(frame, basis="lemma", lowercase=lowercase)

    collocation_rows: list[dict[str, object]] = []
    selected_collocations = _cluster_collocation_table(tables)
    for _, table_row in selected_collocations.iterrows():
        ngram_size = int(table_row["ngram_size"])
        ngram = str(table_row["ngram"])
        occurrence_ids = occurrence_id_index.get(("lemma", ngram_size, ngram), ())
        collocation_rows.append(
            _simple_collocation_row(
                table_row,
                aggregate=aggregate,
                occurrence_ids=occurrence_ids,
                occurrence_infos=occurrence_infos,
            )
        )

    collocations = pd.DataFrame(collocation_rows)
    if not collocations.empty:
        collocations = collocations[
            [
                "collocation",
                "frequency",
                "part_1",
                "part_1_semclasses",
                "part_2",
                "part_2_semclasses",
                "part_3",
                "part_3_semclasses",
                "cluster_counts",
                "intersecting_cluster_count",
            ]
        ]
    pairs = pd.DataFrame()
    examples = pd.DataFrame()

    overview = _collocation_cluster_overview(
        run,
        run_index,
        collocations,
        pairs,
        examples,
    )
    return {
        "overview": overview,
        "collocations": collocations,
        "pairs": pairs,
        "examples": examples,
    }


def _weighted_rate(frame: pd.DataFrame, count_column: str, total_column: str) -> float:
    if frame.empty:
        return float("nan")
    total = int(frame[total_column].sum())
    count = int(frame[count_column].sum())
    return _mean_or_nan(float(count), total)


def _collocation_cluster_overview(
    run: dict[str, object],
    run_index: int,
    collocations: pd.DataFrame,
    pairs: pd.DataFrame,
    examples: pd.DataFrame,
) -> dict[str, object]:
    config = run.get("config", {})
    config = config if isinstance(config, dict) else {}
    return {
        "run_index": run_index,
        "run_label": _run_label(run_index, run),
        "algorithm": config.get("algorithm"),
        "n_clusters": config.get("n_clusters"),
        "collocation_count": int(len(collocations)),
    }


def write_collocation_cluster_excel(
    run_tables: Sequence[dict[str, pd.DataFrame | dict[str, object]]],
    output_path: str | Path,
) -> str:
    """Write collocation cluster-overlap tables to an Excel workbook."""

    try:
        from openpyxl.styles import Alignment
    except ImportError as exc:
        raise ImportError(
            "Excel export requires openpyxl. Install it in Colab with: "
            "pip install openpyxl"
        ) from exc

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    overview = pd.DataFrame([tables["overview"] for tables in run_tables])
    used_sheets = {"overview"}
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        overview.to_excel(writer, sheet_name="overview", index=False)
        for tables in run_tables:
            overview_row = tables["overview"]
            assert isinstance(overview_row, dict)
            run_label = str(overview_row["run_label"])
            table = tables["collocations"]
            assert isinstance(table, pd.DataFrame)
            sheet_name = _sheet_name(f"collocations_{run_label}", used_sheets)
            table.to_excel(writer, sheet_name=sheet_name, index=False)
        for worksheet in writer.sheets.values():
            for row in worksheet.iter_rows(min_row=1, max_row=1):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            worksheet.freeze_panes = "A2"
    return str(output_path)


def write_collocation_cluster_plots(
    run_tables: Sequence[dict[str, pd.DataFrame | dict[str, object]]],
    output_path: str | Path,
) -> list[str]:
    """Write optional collocation cluster-overlap plots if matplotlib is present."""

    try:
        import matplotlib.pyplot as plt
    except ImportError:
        return []

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    plot_paths: list[str] = []
    for tables in run_tables:
        overview = tables["overview"]
        collocations = tables["collocations"]
        pairs = tables["pairs"]
        if not isinstance(overview, dict):
            continue
        run_label = str(overview["run_label"])
        if (
            isinstance(collocations, pd.DataFrame)
            and not collocations.empty
            and {
                "exact_same_cluster_pair_rate",
                "aggregate_same_cluster_pair_rate",
            }.issubset(collocations.columns)
        ):
            scatter_path = output_path.with_name(
                f"{output_path.stem}_{run_label}_exact_vs_aggregate.png"
            )
            fig, axis = plt.subplots(figsize=(7, 5))
            axis.scatter(
                collocations["exact_same_cluster_pair_rate"],
                collocations["aggregate_same_cluster_pair_rate"],
                alpha=0.35,
                s=12,
            )
            axis.set_xlabel("Exact occurrence same-cluster pair rate")
            axis.set_ylabel("Aggregate same-cluster pair rate")
            axis.set_title(f"Collocation cluster overlap: {run_label}")
            fig.tight_layout()
            fig.savefig(scatter_path, dpi=160)
            plt.close(fig)
            plot_paths.append(str(scatter_path))
        if (
            isinstance(pairs, pd.DataFrame)
            and not pairs.empty
            and "aggregate_same_cluster_rate" in pairs.columns
        ):
            histogram_path = output_path.with_name(
                f"{output_path.stem}_{run_label}_pair_rate_hist.png"
            )
            fig, axis = plt.subplots(figsize=(7, 5))
            pairs["aggregate_same_cluster_rate"].dropna().plot.hist(
                bins=20,
                ax=axis,
            )
            axis.set_xlabel("Aggregate pair same-cluster rate")
            axis.set_title(f"Collocation pair cluster-rate distribution: {run_label}")
            fig.tight_layout()
            fig.savefig(histogram_path, dpi=160)
            plt.close(fig)
            plot_paths.append(str(histogram_path))
    return plot_paths


def _attach_plot_paths(
    run_tables: Sequence[dict[str, pd.DataFrame | dict[str, object]]],
    plot_paths: Sequence[str],
) -> None:
    for tables in run_tables:
        overview = tables["overview"]
        if not isinstance(overview, dict):
            continue
        run_label = str(overview["run_label"])
        run_plot_paths = [
            path
            for path in plot_paths
            if f"_{run_label}_" in Path(path).name
        ]
        overview["plot_paths"] = "\n".join(run_plot_paths)


def _default_cluster_output_path(artifact_path: Path) -> Path:
    stem = artifact_path.stem
    if stem.endswith("_artifacts"):
        stem = stem[: -len("_artifacts")]
    return artifact_path.with_name(f"{stem}_{DEFAULT_CLUSTER_OUTPUT_SUFFIX}")


def run_collocation_cluster_experiment(
    config: CollocationClusterConfig | None = None,
    **kwargs: object,
) -> dict[str, object]:
    """Compare CoBaLD collocations with saved clustering artifact runs."""

    if config is None:
        config = CollocationClusterConfig(**kwargs)
    elif kwargs:
        raise ValueError("Pass either config or keyword arguments, not both")
    _validate_cluster_config(config)

    sentences = load_corpus(config.data_dir, splits=config.splits)
    cluster_token_bases: tuple[TokenBasis, ...] = ("lemma",)
    cluster_min_freq = max(config.min_freq, DEFAULT_CLUSTER_MIN_FREQ)
    tables = build_collocation_tables(
        sentences,
        token_bases=cluster_token_bases,
        ngram_sizes=config.ngram_sizes,
        min_freq=cluster_min_freq,
        sort_by=config.sort_by,
        max_rows=config.max_rows,
        lowercase=config.lowercase,
    )
    occurrences = _collocation_occurrences(
        sentences,
        tables,
        token_bases=cluster_token_bases,
        ngram_sizes=config.ngram_sizes,
        lowercase=config.lowercase,
    )
    occurrence_id_index = _occurrence_id_index(occurrences)

    artifact_path = Path(config.artifact_path)
    artifact = load_clustering_artifact(artifact_path)
    runs = artifact["runs"]
    run_tables = [
        analyze_collocation_cluster_run(
            tables,
            occurrences,
            run,
            run_index=run_index,
            lowercase=config.lowercase,
            max_examples=config.max_examples_per_run,
            occurrence_id_index=occurrence_id_index,
        )
        for run_index, run in enumerate(runs)
    ]
    output_path = (
        Path(config.output_path)
        if config.output_path is not None
        else _default_cluster_output_path(artifact_path)
    )
    plot_paths = (
        write_collocation_cluster_plots(run_tables, output_path)
        if config.create_plots
        else []
    )
    _attach_plot_paths(run_tables, plot_paths)
    excel_path = write_collocation_cluster_excel(run_tables, output_path)
    return {
        "config": config,
        "artifact": artifact,
        "artifact_path": str(artifact_path),
        "sentences": sentences,
        "tables": tables,
        "occurrences": occurrences,
        "run_tables": run_tables,
        "output_path": excel_path,
        "plot_paths": plot_paths,
    }


def run_collocation_experiment(
    config: CollocationConfig | None = None,
    **kwargs: object,
) -> dict[str, object]:
    """Load CoBaLD data, score collocations, and write an Excel workbook."""

    if config is None:
        config = CollocationConfig(**kwargs)
    elif kwargs:
        raise ValueError("Pass either config or keyword arguments, not both")
    _validate_config(config)

    sentences = load_corpus(config.data_dir, splits=config.splits)
    tables = build_collocation_tables(
        sentences,
        token_bases=config.token_bases,
        ngram_sizes=config.ngram_sizes,
        min_freq=config.min_freq,
        sort_by=config.sort_by,
        max_rows=config.max_rows,
        lowercase=config.lowercase,
    )
    metadata = _metadata_table(config=config, sentences=sentences, tables=tables)
    output_path = write_collocation_excel(
        tables,
        config.output_path,
        metadata=metadata,
    )
    return {
        "config": config,
        "sentences": sentences,
        "tables": tables,
        "metadata": metadata,
        "output_path": output_path,
    }
