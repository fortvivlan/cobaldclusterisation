from pathlib import Path
import pickle

import pandas as pd
import pytest
from openpyxl import load_workbook

pytest.importorskip("nltk")

from cobaldclusterisation.collocations import (
    _occurrence_id_index,
    build_collocation_tables,
    collocation_table_from_documents,
    run_collocation_cluster_experiment,
    sentence_token_documents,
    write_collocation_excel,
)
from cobaldclusterisation.data import Sentence, Token


def _token(
    form: str,
    lemma: str,
    *,
    token_id: str = "1",
    upos: str = "NOUN",
) -> Token:
    return Token.from_columns(
        [
            token_id,
            form,
            lemma,
            upos,
            "_",
            "_",
            "0",
            "root",
            "_",
            "_",
            "_",
            "_",
        ]
    )


def test_sentence_token_documents_do_not_cross_sentence_boundaries() -> None:
    sentences = [
        Sentence(
            tokens=[
                _token("Нью", "нью", token_id="1"),
                _token("Йорк", "йорк", token_id="2"),
            ]
        ),
        Sentence(
            tokens=[
                _token("Сити", "сити", token_id="1"),
                _token("Холл", "холл", token_id="2"),
            ]
        ),
    ]

    documents = sentence_token_documents(sentences, basis="lemma")
    table = collocation_table_from_documents(
        documents,
        ngram_size=2,
        min_freq=1,
    )

    assert "йорк сити" not in set(table["ngram"])
    assert {"нью йорк", "сити холл"}.issubset(set(table["ngram"]))


def test_build_collocation_tables_separates_lemmas_and_surface_forms() -> None:
    sentences = [
        Sentence(
            tokens=[
                _token("Белые", "белый", token_id="1", upos="ADJ"),
                _token("дома", "дом", token_id="2"),
            ]
        ),
        Sentence(
            tokens=[
                _token("белого", "белый", token_id="1", upos="ADJ"),
                _token("дома", "дом", token_id="2"),
            ]
        ),
    ]

    tables = build_collocation_tables(
        sentences,
        token_bases=("lemma", "form"),
        ngram_sizes=(2,),
        min_freq=2,
    )

    assert tables["lemma_bigrams"].loc[0, "ngram"] == "белый дом"
    assert tables["lemma_bigrams"].loc[0, "frequency"] == 2
    assert tables["form_bigrams"].empty


def test_collocation_table_applies_min_freq_and_sorts_by_pmi() -> None:
    documents = [
        ["a", "b"],
        ["a", "b"],
        ["a", "c"],
        ["x", "y"],
        ["x", "y"],
        ["x", "y"],
    ]

    table = collocation_table_from_documents(
        documents,
        ngram_size=2,
        min_freq=2,
        sort_by="pmi",
    )

    assert set(table["ngram"]) == {"a b", "x y"}
    assert "a c" not in set(table["ngram"])
    assert table["pmi"].is_monotonic_decreasing


def test_write_collocation_excel_creates_expected_sheets(tmp_path: Path) -> None:
    output_path = tmp_path / "collocations.xlsx"
    tables = {
        "lemma_bigrams": pd.DataFrame(
            {
                "token_basis": ["lemma"],
                "ngram_size": [2],
                "ngram": ["белый дом"],
                "frequency": [2],
                "pmi": [1.0],
            }
        )
    }
    metadata = pd.DataFrame({"key": ["min_freq"], "value": [2]})

    write_collocation_excel(tables, output_path, metadata=metadata)

    workbook = load_workbook(output_path)

    assert set(workbook.sheetnames) == {"lemma_bigrams", "metadata"}
    assert workbook["lemma_bigrams"].freeze_panes == "A2"
    assert workbook["metadata"].freeze_panes == "A2"


def test_occurrence_id_index_groups_exact_matches() -> None:
    occurrences = pd.DataFrame(
        {
            "token_basis": ["lemma", "lemma", "form"],
            "ngram_size": [2, 2, 3],
            "ngram": ["белый дом", "белый дом", "new york city"],
            "occurrence_id": [10, 11, 12],
        }
    )

    index = _occurrence_id_index(occurrences)

    assert index[("lemma", 2, "белый дом")] == [10, 11]
    assert index[("form", 3, "new york city")] == [12]


def _write_corpus(path: Path) -> None:
    path.write_text(
        "\n".join(
            [
                "# sent_id = s0",
                "# text = Грозит смертельная опасность.",
                "1\tГрозит\tгрозить\tVERB\t_\t_\t0\troot\t_\t_\t_\tACTION",
                "2\tсмертельная\tсмертельный\tADJ\t_\t_\t1\tamod\t_\t_\t_\tQUALITY",
                "3\tопасность\tопасность\tNOUN\t_\t_\t1\tobj\t_\t_\t_\tDANGER",
                "4\t.\t.\tPUNCT\t_\t_\t1\tpunct\t_\t_\t_\t_",
                "",
                "# sent_id = s1",
                "# text = Грозила смертельная опасность.",
                "1\tГрозила\tгрозить\tVERB\t_\t_\t0\troot\t_\t_\t_\tACTION",
                "2\tсмертельная\tсмертельный\tADJ\t_\t_\t1\tamod\t_\t_\t_\tQUALITY",
                "3\tопасность\tопасность\tNOUN\t_\t_\t1\tobj\t_\t_\t_\tDANGER",
                "4\t.\t.\tPUNCT\t_\t_\t1\tpunct\t_\t_\t_\t_",
                "",
                "# sent_id = s2",
                "# text = Карета скорая помощь.",
                "1\tКарета\tкарета\tNOUN\t_\t_\t0\troot\t_\t_\t_\tVEHICLE",
                "2\tскорая\tскорый\tADJ\t_\t_\t3\tamod\t_\t_\t_\tQUALITY",
                "3\tпомощь\tпомощь\tNOUN\t_\t_\t1\tappos\t_\t_\t_\tHELP",
                "4\t.\t.\tPUNCT\t_\t_\t1\tpunct\t_\t_\t_\t_",
                "",
            ]
        ),
        encoding="utf-8",
    )


def _artifact_tokens(clusters: list[int]) -> pd.DataFrame:
    rows = [
        ("Грозит", "грозить", "VERB", "ACTION", 0, 0, "s0", "Грозит смертельная опасность."),
        (
            "смертельная",
            "смертельный",
            "ADJ",
            "QUALITY",
            0,
            1,
            "s0",
            "Грозит смертельная опасность.",
        ),
        ("опасность", "опасность", "NOUN", "DANGER", 0, 2, "s0", "Грозит смертельная опасность."),
        ("Грозила", "грозить", "VERB", "ACTION", 1, 0, "s1", "Грозила смертельная опасность."),
        (
            "смертельная",
            "смертельный",
            "ADJ",
            "QUALITY",
            1,
            1,
            "s1",
            "Грозила смертельная опасность.",
        ),
        ("опасность", "опасность", "NOUN", "DANGER", 1, 2, "s1", "Грозила смертельная опасность."),
        ("Карета", "карета", "NOUN", "VEHICLE", 2, 0, "s2", "Карета скорая помощь."),
        ("скорая", "скорый", "ADJ", "QUALITY", 2, 1, "s2", "Карета скорая помощь."),
        ("помощь", "помощь", "NOUN", "HELP", 2, 2, "s2", "Карета скорая помощь."),
    ]
    frame = pd.DataFrame(
        [
            {
                "ID": str(token_index + 1),
                "FORM": form,
                "LEMMA": lemma,
                "UPOS": upos,
                "XPOS": "_",
                "FEATS": "_",
                "HEAD": "0",
                "DEPREL": "root",
                "DEPS": "_",
                "MISC": "_",
                "DEEPSLOT": "_",
                "SEMCLASS": semclass,
                "split": "train",
                "sentence_index": sentence_index,
                "token_index": token_index,
                "sentence_id": sentence_id,
                "context_text": context,
            }
            for form, lemma, upos, semclass, sentence_index, token_index, sentence_id, context in rows
        ]
    )
    frame.insert(0, "cluster", clusters)
    frame["AUTO_SEMCLASS"] = [f"cluster_{cluster}" for cluster in clusters]
    return frame


def _write_artifact(path: Path) -> None:
    runs = [
        {
            "config": {"algorithm": "minibatch_kmeans", "n_clusters": 100},
            "labels": [0, 0, 1, 0, 0, 1, 2, 3, 4],
            "annotated_tokens": _artifact_tokens([0, 0, 1, 0, 0, 1, 2, 3, 4]),
        },
        {
            "config": {"algorithm": "minibatch_kmeans", "n_clusters": 200},
            "labels": [0, 1, 1, 0, 1, 1, 2, 2, 3],
            "annotated_tokens": _artifact_tokens([0, 1, 1, 0, 1, 1, 2, 2, 3]),
        },
    ]
    with path.open("wb") as handle:
        pickle.dump({"runs": runs, "metadata": {}}, handle)


def test_run_collocation_cluster_experiment_writes_overlap_workbook(
    tmp_path: Path,
) -> None:
    _write_corpus(tmp_path / "train.conllu")
    results_dir = tmp_path / "results"
    results_dir.mkdir()
    artifact_path = results_dir / "toy_artifacts.pkl"
    _write_artifact(artifact_path)

    result = run_collocation_cluster_experiment(
        data_dir=tmp_path,
        splits=("train",),
        token_bases=("lemma",),
        ngram_sizes=(2, 3),
        min_freq=1,
        artifact_path=artifact_path,
        create_plots=False,
    )

    assert Path(result["output_path"]).parent == results_dir
    workbook = load_workbook(result["output_path"])
    assert {
        "overview",
        "collocations_k100",
        "pairs_k100",
        "examples_k100",
        "collocations_k200",
        "pairs_k200",
        "examples_k200",
    }.issubset(set(workbook.sheetnames))
    assert workbook["overview"].freeze_panes == "A2"

    run_tables = result["run_tables"]
    first_run = run_tables[0]
    assert isinstance(first_run["pairs"], pd.DataFrame)
    pair = first_run["pairs"].loc[
        (first_run["pairs"]["ngram"] == "грозить смертельный")
        & (first_run["pairs"]["part_a"] == "грозить")
        & (first_run["pairs"]["part_b"] == "смертельный")
    ].iloc[0]
    assert pair["exact_same_cluster_rate"] == pytest.approx(1.0)
    assert pair["aggregate_same_cluster_rate"] == pytest.approx(1.0)
    assert "ACTION" in pair["part_a_semclasses"]
    assert "QUALITY" in pair["part_b_semclasses"]

    trigram = first_run["collocations"].loc[
        first_run["collocations"]["ngram"] == "грозить смертельный опасность"
    ].iloc[0]
    assert trigram["exact_all_parts_same_cluster_rate"] == pytest.approx(0.0)
    assert trigram["exact_same_cluster_pair_rate"] == pytest.approx(1 / 3)

    second_run = run_tables[1]
    assert isinstance(second_run["pairs"], pd.DataFrame)
    split_pair = second_run["pairs"].loc[
        (second_run["pairs"]["ngram"] == "грозить смертельный")
        & (second_run["pairs"]["part_a"] == "грозить")
        & (second_run["pairs"]["part_b"] == "смертельный")
    ].iloc[0]
    assert split_pair["exact_same_cluster_rate"] == pytest.approx(0.0)
