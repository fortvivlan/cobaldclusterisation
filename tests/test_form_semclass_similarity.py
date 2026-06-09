from pathlib import Path
import pickle

import numpy as np
from openpyxl import load_workbook
import pandas as pd
import pytest

from cobaldclusterisation.form_semclass_similarity import (
    analyze_run,
    infer_embeddings_path,
    run_form_semclass_similarity_experiment,
    validate_alignment,
)


def _tokens() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "ID": ["1", "1", "1", "1", "1"],
            "FORM": ["Идет", "идет", "идет", "идет", "дом"],
            "LEMMA": ["идти", "идти", "идти", "идти", "дом"],
            "UPOS": ["VERB", "VERB", "VERB", "VERB", "NOUN"],
            "XPOS": ["_"] * 5,
            "FEATS": ["_"] * 5,
            "HEAD": ["0"] * 5,
            "DEPREL": ["root"] * 5,
            "DEPS": ["_"] * 5,
            "MISC": ["_"] * 5,
            "DEEPSLOT": ["_"] * 5,
            "SEMCLASS": ["MOVE", "MOVE", "RAIN", "RAIN", "BUILDING"],
            "split": ["train"] * 5,
            "sentence_index": [0, 1, 2, 3, 4],
            "token_index": [0, 0, 0, 0, 0],
            "sentence_id": ["s0", "s1", "s2", "s3", "s4"],
            "context_text": [
                "Вася идет.",
                "Петя идет.",
                "Дождь идет.",
                "Снег идет.",
                "дом стоит.",
            ],
            "embedding_index": [0, 1, 2, 3, 4],
        }
    )


def _embeddings() -> np.ndarray:
    return np.asarray(
        [
            [1.0, 0.0],
            [1.0, 0.0],
            [0.0, 1.0],
            [0.0, 1.0],
            [0.0, 1.0],
        ],
        dtype=np.float32,
    )


def _run() -> dict[str, object]:
    annotated = _tokens().copy()
    annotated.insert(0, "cluster", [0, 1, 0, 0, 2])
    annotated["AUTO_SEMCLASS"] = [
        "cluster_0",
        "cluster_1",
        "cluster_0",
        "cluster_0",
        "cluster_2",
    ]
    return {
        "config": {"algorithm": "minibatch_kmeans", "n_clusters": 100},
        "labels": np.asarray([0, 1, 0, 0, 2], dtype=np.int64),
        "annotated_tokens": annotated,
    }


def test_analyze_run_calculates_exact_cosine_and_cluster_rates() -> None:
    result = analyze_run(
        _tokens(),
        _embeddings(),
        _run(),
        run_index=0,
        max_examples=10,
        examples_per_form=2,
    )

    forms = result["forms"]
    assert isinstance(forms, pd.DataFrame)
    row = forms.loc[forms["form_lower"] == "идет"].iloc[0]

    assert row["diff_semclass_pair_count"] == 4
    assert row["diff_semclass_mean_cosine"] == pytest.approx(0.0)
    assert row["diff_semclass_same_cluster_rate"] == pytest.approx(0.5)
    assert row["same_semclass_pair_count"] == 2
    assert row["same_semclass_mean_cosine"] == pytest.approx(1.0)
    assert row["same_semclass_different_cluster_rate"] == pytest.approx(0.5)
    assert row["same_minus_diff_cosine"] == pytest.approx(1.0)

    overview = result["overview"]
    assert isinstance(overview, dict)
    assert overview["diff_semclass_same_cluster_rate"] == pytest.approx(0.5)
    assert overview["same_semclass_different_cluster_rate"] == pytest.approx(0.5)


def test_validate_alignment_rejects_mismatched_tokens() -> None:
    embedding_tokens = _tokens()
    run = _run()
    run["annotated_tokens"].loc[0, "FORM"] = "другой"

    with pytest.raises(ValueError, match="FORM"):
        validate_alignment(embedding_tokens, _embeddings(), [run])


def test_infer_embeddings_path_uses_drive_root_default(tmp_path: Path) -> None:
    artifact_path = (
        tmp_path
        / "drive"
        / "MyDrive"
        / "cobald_outputs"
        / "results"
        / "run_artifacts.pkl"
    )
    artifact_path.parent.mkdir(parents=True)
    embeddings_path = artifact_path.parent.parent / "rubert_tiny2_cobald.pkl"
    embeddings_path.write_bytes(b"pickle")

    assert infer_embeddings_path(artifact_path, {"metadata": {}}) == embeddings_path


def test_run_experiment_writes_expected_workbook(tmp_path: Path) -> None:
    root = tmp_path / "drive" / "MyDrive" / "cobald_outputs"
    results_dir = root / "results"
    results_dir.mkdir(parents=True)
    artifact_path = results_dir / "run_artifacts.pkl"
    embeddings_path = root / "rubert_tiny2_cobald.pkl"
    output_path = tmp_path / "summary.xlsx"

    with artifact_path.open("wb") as handle:
        pickle.dump({"runs": [_run()], "metadata": {}}, handle)
    with embeddings_path.open("wb") as handle:
        pickle.dump({"embeddings": _embeddings(), "tokens": _tokens()}, handle)

    result = run_form_semclass_similarity_experiment(
        artifact_path=artifact_path,
        output_path=output_path,
        create_plots=False,
    )

    workbook = load_workbook(result["output_path"])
    assert {
        "overview",
        "forms_k100",
        "classpairs_k100",
        "within_k100",
        "examples_k100",
    }.issubset(set(workbook.sheetnames))
    assert workbook["overview"].freeze_panes == "A2"
